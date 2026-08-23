#!/usr/bin/env python3
"""Optional xlsx export. Derived file — do not treat as the spine.

Run: python3 scripts/build_xlsx.py
Out: data/Jeff Story Master Song Catalog.xlsx (gitignored-style derived; not required)
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
cat = json.load(open(os.path.join(HERE, "data", "master_catalog.json")))
wb = Workbook()

HDR = Font(name='Arial', bold=True, color='FFFFFF', size=10)
HFILL = PatternFill('solid', fgColor='1F3864')
BODY = Font(name='Arial', size=10)
WRAP = Alignment(wrap_text=True, vertical='top')
TIER = {'high': PatternFill('solid', fgColor='C6EFCE'), 'medium': PatternFill('solid', fgColor='FFEB9C')}

def sheet(ws, cols, rows, widths):
    for j,(h,_) in enumerate(cols,1):
        c=ws.cell(1,j,h); c.font=HDR; c.fill=HFILL; c.alignment=WRAP
        ws.column_dimensions[get_column_letter(j)].width=widths[j-1]
    for i,r in enumerate(rows,2):
        for j,(_,k) in enumerate(cols,1):
            v=r.get(k,'')
            if isinstance(v,list): v='; '.join(map(str,v))
            c=ws.cell(i,j,v); c.font=BODY; c.alignment=WRAP
    ws.freeze_panes='A2'
    ws.auto_filter.ref=f"A1:{get_column_letter(len(cols))}{len(rows)+1}"

# --- Songs sheet ---
ws = wb.active; ws.title='Songs'
cols=[('ID','song_id'),('Canonical Title','canonical_title'),('Artist/Project','artist_project'),
 ('Class','classification'),('Alt Titles','alt_titles'),('Writers','writers'),('Rights','rights_confidence'),
 ('Stage','stage'),('Key','key'),('BPM','bpm'),('Potential /100','potential'),('Readiness /100','readiness'),
 ('Confidence','confidence'),('Lyric Status','lyric_status'),('Audio Status','audio_status'),
 ('Theme','theme'),('Hook','hook'),('Sources','sources'),('Best Source','best_source'),
 ('Next Action','next_action'),('Open Questions','open_questions'),('Evidence / Notes','notes')]
rows=sorted(cat['songs'], key=lambda s:(-(s['potential'] or 0), s['song_id']))
sheet(ws, cols, rows, [9,26,20,14,20,20,16,22,12,8,10,10,12,16,13,30,30,45,24,35,30,50])
for i,r in enumerate(rows,2):
    f=TIER.get(r['confidence'])
    if f: ws.cell(i,13).fill=f

# --- Active WIP ---
ws2=wb.create_sheet('Active WIP')
wip=[dict(slot='Flagship', id='ST-0001', title='Turn Over The Flag',
  low='Listen to mix 1.6 once on a phone speaker; note anything that bugs you (5 min)',
  normal='Record the bridge "oohs" (1 part, 3-4 layered takes)',
  deep='Guitar doubles + octave line, rough balance, bounce mix 1.7'),
 dict(slot='Quick win', id='ST-0004', title='Manic',
  low='Listen to bounce 4. manic v1.4 once (3 min)',
  normal='Track the chorus guitar octave part (the ONLY remaining item)',
  deep='Comp + tune the part, drop into the mix, bounce final'),
 dict(slot='Experimental', id='ST-0009', title='Long Long Drive',
  low='Re-read the lyric doc; circle the one line that hits hardest (3 min)',
  normal='Rough-track the documented background-vox outro over the 2019 recording',
  deep='Full-band arrangement experiment from the Am/123 chart (Suno test pending approval)')]
cols2=[('Slot','slot'),('ID','id'),('Song','title'),('Very-low-energy action','low'),
 ('Focused-session action','normal'),('Deep production action','deep')]
sheet(ws2, cols2, wip, [12,9,20,40,40,40])

# --- Covers ---
ws3=wb.create_sheet('Covers & Reference')
cols3=[('Title','title'),('Original Artist','original_artist'),('Context','context'),('Class','classification')]
sheet(ws3, cols3, sorted(cat['covers'], key=lambda c:(c['original_artist'],c['title'])), [30,26,45,10])

# --- Read Me ---
ws4=wb.create_sheet('Read Me')
notes=[
 ('Jeff Story Master Song Catalog — v1, generated 2026-08-18 (Session 1)',),
 ('',),
 ('HONESTY NOTE: No audio has been listened to in this session. Every audio item = metadata only.',),
 ('Scores exist ONLY for songs whose lyrics were actually read (14 songs). Blank = insufficient evidence, not zero.',),
 ('',),
 ('Potential /100 = hook 25 + lyric 20 + emotional truth 20 + structure 15 + identity fit 10 + replay/live 10',),
 ('Readiness /100 = completeness 25 + clarity 20 + source usability 15 + manageable remaining work 25 + feasibility 15',),
 ('These are deliberately separate; a song can be high-potential/low-readiness and vice versa.',),
 ('',),
 ('Confidence colors on Songs sheet: green = high, yellow = medium, none = low/insufficient.',),
 ('IDs: ST = Stalemate (ST-01xx = 2010 album), SD = Something Dirty, JS = solo, UNK = authorship uncertain.',),
 ('Covers are cataloged on their own sheet and are NEVER ranked against originals.',),
 ('Source of truth for resume: claude.ai project "2026 Song Organization" > claude/ docs.',),
]
for i,row in enumerate(notes,1):
    c=ws4.cell(i,1,row[0]); c.font=Font(name='Arial', size=10, bold=(i==1))
ws4.column_dimensions['A'].width=110

out = os.path.join(HERE, "data", "Jeff Story Master Song Catalog.xlsx")
wb.save(out)
print('saved', out)
